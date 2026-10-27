from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
NAVY="1F2A44"; hdr=Font(bold=True,color="FFFFFF"); hfill=PatternFill("solid",fgColor=NAVY)
inp=PatternFill("solid",fgColor="FFF8DC"); gry=PatternFill("solid",fgColor="F2F4F7")
thin=Side(style="thin",color="D0D5DD"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); s=wb.active; s.title="Start here"
s.column_dimensions["A"].width=3; s.column_dimensions["B"].width=95
txt=[("Options Risk & Position Size Calculator",18,True),
("Work out the maximum loss, breakeven and position size of a trade before you place it.",11,False),("",11,False),
("How to use it",13,True),
("1. 'Account' tab: type your account size and the percent of it you allow yourself to lose on one trade. Yellow cells are inputs.",11,False),
("2. 'Long option' tab: enter the premium and contracts for a single call or put.",11,False),
("3. 'Vertical spread' tab: enter strikes and the net debit or credit for a debit or credit spread.",11,False),
("4. Each tab tells you whether the trade fits your risk limit and the most contracts that would.",11,False),
("5. Prices are per share, as quoted (an option quoted at 1.25 costs $125 per contract). Multiplier is 100.",11,False),("",11,False),
("Please read",13,True),
("This is arithmetic, not advice or a signal. It uses the numbers you type and cannot know fills, slippage, assignment or early-exercise risk.",11,False),
("Max-loss figures assume you hold to expiry or exit at zero; real exits are usually better, but gaps can be worse for shares you end up holding.",11,False),
("Trading options involves risk of loss. Built with AI assistance (Claude) and tested by recalculating every formula.",11,False)]
for i,(t,z,b) in enumerate(txt,start=2):
    c=s.cell(row=i,column=2,value=t); c.font=Font(size=z,bold=b,color=NAVY if b else "000000"); c.alignment=Alignment(wrap_text=True)
def sheet(name):
    w=wb.create_sheet(name); w.column_dimensions["A"].width=36; w.column_dimensions["B"].width=16; w.column_dimensions["C"].width=62; return w
def put(w,r,label,val,nf=None,inputcell=False,note=""):
    w.cell(row=r,column=1,value=label).font=Font(bold=True)
    c=w.cell(row=r,column=2,value=val); c.border=box; c.fill=inp if inputcell else gry
    if nf: c.number_format=nf
    w.cell(row=r,column=3,value=note).font=Font(italic=True,color="667085")
def title(w,t):
    w["A1"]=t; w["A1"].font=Font(size=15,bold=True,color=NAVY)
USD='$#,##0.00'; 
a=sheet("Account"); title(a,"Your account and risk limit")
put(a,3,"Account size ($)",5000,USD,True)
put(a,4,"Max loss per trade (% of account)",0.05,"0.0%",True,"Example: 5% means you will not risk more than 5% on one trade")
put(a,5,"Max loss per trade ($)","=B3*B4",USD,note="Account size times the percent")
put(a,6,"Max loss per day (% of account)",0.10,"0.0%",True,"Optional: your stop-for-the-day level")
put(a,7,"Max loss per day ($)","=B3*B6",USD)
put(a,8,"Trades you could lose in a row before the daily limit","=IF(B5=0,\"\",ROUNDDOWN(B7/B5,0))","0",note="If every trade lost its full max")
a["A10"]="Position size for any trade"; a["A10"].font=Font(bold=True,color=NAVY,size=12)
put(a,11,"Max loss per contract ($)",50,USD,True,"From the other tabs, or type your own")
put(a,12,"Contracts allowed","=IF(B11<=0,\"\",ROUNDDOWN(B5/B11,0))","0",note="Whole contracts only")
put(a,13,"Result",'=IF(B12="","",IF(B12=0,"Too big for your limit: even 1 contract risks more than "&TEXT(B5,"$#,##0.00"),B12&" contract(s) fit"))')
# long option
l=sheet("Long option"); title(l,"Long call or put")
put(l,3,"Premium paid per share ($)",1.25,'0.00',True,"The price you pay, as quoted")
put(l,4,"Contracts",2,"0",True)
put(l,5,"Multiplier",100,"0",True,"100 for US stock options")
put(l,6,"Fees, round trip ($)",2,USD,True)
put(l,7,"Call or put",'Call',None,True)
put(l,8,"Strike",500,'0.00',True)
dv=DataValidation(type="list",formula1='"Call,Put"'); l.add_data_validation(dv); dv.add("B7")
l["A10"]="Results"; l["A10"].font=Font(bold=True,color=NAVY,size=12)
put(l,11,"Cost / max loss ($)","=B3*B4*B5+B6",USD,note="You can lose the premium plus fees, no more")
put(l,12,"Max loss per contract ($)","=B3*B5+IF(B4=0,0,B6/B4)",USD)
put(l,13,"Breakeven at expiry","=IF(B7=\"Call\",B8+B3+IF(B4*B5=0,0,B6/(B4*B5)),B8-B3-IF(B4*B5=0,0,B6/(B4*B5)))",'0.00')
put(l,14,"Max profit","=IF(B7=\"Call\",\"Unlimited\",(B8-B3)*B4*B5-B6)",USD,note="Put: strike minus premium, if the stock goes to zero")
put(l,15,"Move needed to break even (%)","=IF(B8=0,\"\",ABS(B13-B8)/B8)","0.00%",note="From the strike, as a share of the strike")
put(l,16,"Fits your risk limit?",'=IF(B11<=Account!B5,"Yes","No: max loss "&TEXT(B11,"$#,##0.00")&" is above your limit of "&TEXT(Account!B5,"$#,##0.00"))')
put(l,17,"Most contracts that fit","=IF(B12<=0,\"\",ROUNDDOWN(Account!B5/B12,0))","0")
# vertical
v=sheet("Vertical spread"); title(v,"Vertical spread (debit or credit)")
put(v,3,"Type",'Bull call (debit)',None,True,"Pick from the list")
dv2=DataValidation(type="list",formula1='"Bull call (debit),Bear put (debit),Bull put (credit),Bear call (credit)"'); v.add_data_validation(dv2); dv2.add("B3")
put(v,4,"Lower strike",500,'0.00',True)
put(v,5,"Higher strike",505,'0.00',True)
put(v,6,"Net price per share ($)",2.10,'0.00',True,"Debit you pay, or credit you receive, as quoted")
put(v,7,"Contracts",1,"0",True)
put(v,8,"Multiplier",100,"0",True)
put(v,9,"Fees, round trip ($)",2,USD,True)
v["A11"]="Results"; v["A11"].font=Font(bold=True,color=NAVY,size=12)
put(v,12,"Width ($ per share)","=B5-B4",'0.00')
put(v,13,"Debit or credit?",'=IF(RIGHT(B3,7)="(debit)","Debit","Credit")')
put(v,14,"Max loss ($)",'=IF(B13="Debit",B6*B7*B8+B9,(B12-B6)*B7*B8+B9)',USD,note="Debit: what you pay. Credit: width minus credit")
put(v,15,"Max profit ($)",'=IF(B13="Debit",(B12-B6)*B7*B8-B9,B6*B7*B8-B9)',USD)
put(v,16,"Breakeven at expiry",'=IF(B3="Bull call (debit)",B4+B6,IF(B3="Bear put (debit)",B5-B6,IF(B3="Bull put (credit)",B5-B6,B4+B6)))','0.00',note="Before fees")
put(v,17,"Reward : risk",'=IF(B14=0,"",B15/B14)','0.00')
put(v,18,"Max loss per contract ($)",'=IF(B7=0,"",B14/B7)',USD)
put(v,19,"Fits your risk limit?",'=IF(B14<=Account!B5,"Yes","No: max loss "&TEXT(B14,"$#,##0.00")&" is above your limit of "&TEXT(Account!B5,"$#,##0.00"))')
put(v,20,"Most contracts that fit",'=IF(N(B18)<=0,"",ROUNDDOWN(Account!B5/B18,0))',"0")
put(v,21,"Check on inputs",'=IF(B12<=0,"Higher strike must be above lower strike",IF(AND(B13="Debit",B6>=B12),"Debit is at or above width: no profit possible",IF(AND(B13="Credit",B6>=B12),"Credit is at or above width: check the price","OK")))')
wb.save("../Options Risk and Position Size Calculator.xlsx"); print("ok")
