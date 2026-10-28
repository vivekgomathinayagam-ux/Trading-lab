from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
NAVY="1F2A44"; inp=PatternFill("solid",fgColor="FFF8DC"); gry=PatternFill("solid",fgColor="F2F4F7")
thin=Side(style="thin",color="D0D5DD"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
H=Font(bold=True,color=NAVY,size=14)
wb=Workbook(); s=wb.active; s.title="Start here"; s.column_dimensions["A"].width=3; s.column_dimensions["B"].width=95
lines=[("Options Trading Plan & Checklist Pack",18,True),("A pre-trade checklist, a daily plan and a weekly review, so you follow your own rules.",11,False),("",11,False),
("How to use it",13,True),
("1. 'Pre-Trade Checklist': edit the rules in column B to be YOUR rules, then answer Y or N for each trade. It says GO only if every required rule is a Y.",11,False),
("2. 'Daily Plan': fill it in before the open: your max trades, loss limit, setups you will take and the ones you will skip.",11,False),
("3. 'Weekly Review': log each day, count rule breaks, and write one change for next week. Totals calculate themselves.",11,False),
("Yellow cells are inputs; grey cells are formulas.",11,False),("",11,False),("Please read",13,True),
("This is a planning and record-keeping tool. It is not advice or a signal and does not predict or improve results by itself. The default rules are examples; replace them with your own.",11,False),
("Options trading involves risk of loss. Built with AI assistance (Claude) and tested by recalculating every formula.",11,False)]
for i,(t,z,b) in enumerate(lines,start=2):
    c=s.cell(row=i,column=2,value=t); c.font=Font(size=z,bold=b,color=NAVY if b else "000000"); c.alignment=Alignment(wrap_text=True)

c=wb.create_sheet("Pre-Trade Checklist"); c["A1"]="Pre-trade checklist"; c["A1"].font=H
for col,w in zip("ABCD",(5,70,12,12)): c.column_dimensions[col].width=w
for j,t in enumerate(["#","Rule (edit to your own)","Answer Y/N","Required?"],1):
    x=c.cell(row=3,column=j,value=t); x.font=Font(bold=True,color="FFFFFF"); x.fill=PatternFill("solid",fgColor=NAVY)
rules=[("The setup is one of the setups written in my plan today","Y"),("Max loss of the whole position is within my per-trade limit","Y"),("I priced it at the ask when buying / bid when selling, fees included","Y"),
("My stop or exit level is written down before entry","Y"),("I am under my max trades for today","Y"),("I am under my daily loss limit","Y"),("No major scheduled news in the next 15 minutes that I did not plan for","N"),
("I am calm: not trying to win back a loss","Y"),("The trade is defined-risk (no naked short options)","Y"),("Spread / liquidity is acceptable for this contract","N")]
for i,(r,req) in enumerate(rules,start=4):
    c.cell(row=i,column=1,value=i-3); c.cell(row=i,column=2,value=r).alignment=Alignment(wrap_text=True)
    a=c.cell(row=i,column=3); a.fill=inp; a.border=box; q=c.cell(row=i,column=4,value=req); q.fill=inp; q.border=box
    c.cell(row=i,column=2).fill=inp
last=3+len(rules)
dv=DataValidation(type="list",formula1='"Y,N"',allow_blank=True); c.add_data_validation(dv); dv.add(f"C4:C{last}"); dv.add(f"D4:D{last}")
c.cell(row=last+2,column=2,value="Required rules answered Y").font=Font(bold=True)
c.cell(row=last+2,column=3,value=f'=COUNTIFS(D4:D{last},"Y",C4:C{last},"Y")').fill=gry
c.cell(row=last+3,column=2,value="Required rules in total").font=Font(bold=True)
c.cell(row=last+3,column=3,value=f'=COUNTIF(D4:D{last},"Y")').fill=gry
c.cell(row=last+4,column=2,value="Result").font=Font(bold=True)
res=c.cell(row=last+4,column=3,value=f'=IF(COUNTIFS(D4:D{last},"Y",C4:C{last},"")>0,"INCOMPLETE",IF(C{last+2}=C{last+3},"GO","WAIT"))'); res.fill=gry; res.font=Font(bold=True)
c.cell(row=last+5,column=2,value="Clear column C before each new trade. GO = every required rule is Y. WAIT = at least one required rule is N.").font=Font(italic=True,color="667085")
c.conditional_formatting.add(f"C{last+4}",CellIsRule(operator="equal",formula=['"GO"'],fill=PatternFill("solid",bgColor="C6EFCE")))
c.conditional_formatting.add(f"C{last+4}",CellIsRule(operator="equal",formula=['"WAIT"'],fill=PatternFill("solid",bgColor="FFC7CE")))

d=wb.create_sheet("Daily Plan"); d["A1"]="Daily plan (fill in before the open)"; d["A1"].font=H
d.column_dimensions["A"].width=40; d.column_dimensions["B"].width=60
items=["Date","Account size ($)","Max loss per trade (%)","Max loss per trade ($)","Max loss for the day (%)","Max loss for the day ($)","Max number of trades","Setups I will take today","Setups I will skip today","Events / news to respect","If I hit my daily limit I will","One thing I will do better today"]
for i,t in enumerate(items,start=3):
    d.cell(row=i,column=1,value=t).font=Font(bold=True)
    x=d.cell(row=i,column=2); x.border=box; x.fill=inp; x.alignment=Alignment(wrap_text=True)
d["B4"]=5000; d["B4"].number_format='$#,##0.00'; d["B5"]=0.05; d["B5"].number_format="0.0%"
d["B6"]="=B4*B5"; d["B6"].fill=gry; d["B6"].number_format='$#,##0.00'
d["B7"]=0.10; d["B7"].number_format="0.0%"; d["B8"]="=B4*B7"; d["B8"].fill=gry; d["B8"].number_format='$#,##0.00'; d["B9"]=3
d["B10"]="Example: the setup from my written plan"; 

w=wb.create_sheet("Weekly Review"); w["A1"]="Weekly review"; w["A1"].font=H
heads=["Date","Trades taken","Rule breaks (count)","Day P&L ($)","What I did well","What I would change"]
for j,t in enumerate(heads,1):
    x=w.cell(row=3,column=j,value=t); x.font=Font(bold=True,color="FFFFFF"); x.fill=PatternFill("solid",fgColor=NAVY); x.alignment=Alignment(wrap_text=True)
for col,wd in zip("ABCDEF",(14,14,16,14,40,40)): w.column_dimensions[col].width=wd
for r in range(4,9):
    for j in range(1,7):
        x=w.cell(row=r,column=j); x.fill=inp; x.border=box
    w.cell(row=r,column=4).number_format='$#,##0.00;[Red]-$#,##0.00'
sample=[("2026-01-05",2,0,35),("2026-01-06",3,1,-60),("2026-01-07",1,0,20)]
for r,(dt,n,b,p) in enumerate(sample,start=4):
    w.cell(row=r,column=1,value=dt); w.cell(row=r,column=2,value=n); w.cell(row=r,column=3,value=b); w.cell(row=r,column=4,value=p)
w["A10"]="Totals"; w["A10"].font=Font(bold=True)
for col,f in (("B","=SUM(B4:B8)"),("C","=SUM(C4:C8)"),("D","=SUM(D4:D8)")):
    w[f"{col}10"]=f; w[f"{col}10"].fill=gry; w[f"{col}10"].font=Font(bold=True)
w["D10"].number_format='$#,##0.00;[Red]-$#,##0.00'
w["A11"]="Days with zero rule breaks"; w["B11"]='=COUNTIFS(C4:C8,0,B4:B8,">0")'; w["B11"].fill=gry
w["A12"]="Days traded"; w["B12"]='=COUNTIF(B4:B8,">0")'; w["B12"].fill=gry
w["A13"]="Rule-following rate"; w["B13"]='=IF(B12=0,"",B11/B12)'; w["B13"].number_format="0%"; w["B13"].fill=gry
w["A15"]="The one change I will make next week:"; w["A15"].font=Font(bold=True); w["B15"].fill=inp; w.merge_cells("B15:F15")
w["A17"]="Sample rows above are made up. Overwrite them. The rule-following rate measures whether you followed your plan, not whether trades made money."
w["A17"].font=Font(italic=True,color="667085")
wb.save("Options Trading Plan and Checklist Pack.xlsx")
