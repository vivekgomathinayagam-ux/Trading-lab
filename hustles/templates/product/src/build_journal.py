"""Builds Trade Journal.xlsx (formulas only; nothing hard-coded in stats)."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference

NAVY="1F2A44"; INPUT="FFF8DC"; GREY="F2F4F7"
hdr=Font(bold=True,color="FFFFFF"); hfill=PatternFill("solid",fgColor=NAVY)
inp=PatternFill("solid",fgColor=INPUT); gry=PatternFill("solid",fgColor=GREY)
thin=Side(style="thin",color="D0D5DD"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
N=500  # journal rows

wb=Workbook()
# ---------- Start here
s=wb.active; s.title="Start here"
s.column_dimensions["A"].width=3; s.column_dimensions["B"].width=95
lines=[("Options Day Trader Journal",18,True),
("A plain trade journal that works out your win rate, expectancy, profit factor and drawdown for you.",11,False),("",11,False),
("How to use it",13,True),
("1. Go to the 'Journal' tab. Add one row per trade. Fill only the yellow cells.",11,False),
("2. Planned risk ($) is the most you were willing to lose on the trade BEFORE you entered (for a long option, usually the premium paid).",11,False),
("3. Multiplier is 100 for US stock options, 1 for shares. Fees is your total commission and fees for the round trip.",11,False),
("4. P&L ($), R-multiple, win/loss and the running total fill in by themselves. Do not type in the grey cells.",11,False),
("5. 'Stats' shows the whole picture. 'Calendar' shows each day's P&L for the month you choose. 'By Setup' splits results by setup tag.",11,False),
("6. The Journal has room for 500 trades. To add more, copy the last row's formulas down (every grey cell).",11,False),("",11,False),
("Words used",13,True),
("R-multiple: your P&L divided by the risk you planned. +1R means you made what you risked; -1R means you lost all of it.",11,False),
("Expectancy: the average R (or $) you make per trade. Positive means the sample is profitable after fees.",11,False),
("Profit factor: total winnings divided by total losses. Above 1 means winners outweigh losers.",11,False),
("Max drawdown: the largest fall from a peak in your running P&L.",11,False),("",11,False),
("Please read",13,True),
("This is a record-keeping tool. It does not give trade signals or advice, and keeping a journal does not make anyone profitable.",11,False),
("Trading options involves risk of loss. The sample rows are made-up examples; delete them before you start.",11,False),
("Built with AI assistance (Claude) and tested by recalculating every formula.",11,False)]
for i,(t,sz,b) in enumerate(lines,start=2):
    c=s.cell(row=i,column=2,value=t); c.font=Font(size=sz,bold=b,color=NAVY if b else "000000"); c.alignment=Alignment(wrap_text=True,vertical="top")

# ---------- Journal
j=wb.create_sheet("Journal")
cols=["Date","Symbol","Setup tag","Direction","Entry price","Exit price","Contracts / shares","Multiplier","Fees ($)","Planned risk ($)","Notes",
      "P&L ($)","R-multiple","Result","Running P&L ($)","Peak ($)","Drawdown ($)"]
widths=[12,9,16,10,11,11,11,10,9,12,32,11,11,9,14,11,12]
for i,(c,w) in enumerate(zip(cols,widths),start=1):
    x=j.cell(row=1,column=i,value=c); x.font=hdr; x.fill=hfill; x.alignment=Alignment(wrap_text=True,horizontal="center",vertical="center"); x.border=box
    j.column_dimensions[chr(64+i)].width=w
j.row_dimensions[1].height=32; j.freeze_panes="A2"
sample=[(dt.date(2026,10,5),"SPY","ORB","Long",1.20,1.85,2,100,2.0,240,"SAMPLE: delete me"),
(dt.date(2026,10,5),"QQQ","Reversal","Long",2.10,1.40,1,100,1.0,210,"SAMPLE"),
(dt.date(2026,10,6),"SPY","ORB","Short",1.50,2.30,2,100,2.0,300,"SAMPLE"),
(dt.date(2026,10,7),"SPY","Vertical","Long",0.80,0.55,3,100,3.0,240,"SAMPLE"),
(dt.date(2026,10,8),"QQQ","ORB","Long",1.00,1.90,2,100,2.0,200,"SAMPLE")]
for r in range(2,N+2):
    for c in range(1,12):
        x=j.cell(row=r,column=c); x.fill=inp; x.border=box
    j.cell(row=r,column=1).number_format="yyyy-mm-dd"
    for c in (5,6): j.cell(row=r,column=c).number_format="0.00"
    j.cell(row=r,column=9).number_format="0.00"; j.cell(row=r,column=10).number_format="#,##0.00"
    if r-2<len(sample):
        for c,v in enumerate(sample[r-2],start=1): j.cell(row=r,column=c,value=v)
    # P&L: long = (exit-entry), short = (entry-exit); direction must be Long/Short
    j.cell(row=r,column=12,value=f'=IF(OR(E{r}="",F{r}=""),"",ROUND(IF(D{r}="Short",E{r}-F{r},F{r}-E{r})*G{r}*H{r}-I{r},2))')
    j.cell(row=r,column=13,value=f'=IF(OR(L{r}="",N(J{r})<=0),"",L{r}/J{r})')
    j.cell(row=r,column=14,value=f'=IF(L{r}="","",IF(L{r}>0,"Win",IF(L{r}<0,"Loss","Flat")))')
    j.cell(row=r,column=15,value=f'=IF(L{r}="","",SUM($L$2:L{r}))')
    j.cell(row=r,column=16,value=f'=IF(L{r}="","",MAX(0,MAX($O$2:O{r})))')
    j.cell(row=r,column=17,value=f'=IF(L{r}="","",P{r}-O{r})')
    for c in range(12,18):
        x=j.cell(row=r,column=c); x.fill=gry; x.border=box
    j.cell(row=r,column=12).number_format='#,##0.00;[Red]-#,##0.00'
    j.cell(row=r,column=13).number_format='0.00"R";[Red]-0.00"R"'
    for c in (15,16,17): j.cell(row=r,column=c).number_format='#,##0.00;[Red]-#,##0.00'
dv=DataValidation(type="list",formula1='"Long,Short"',allow_blank=True); j.add_data_validation(dv); dv.add(f"D2:D{N+1}")
j.conditional_formatting.add(f"N2:N{N+1}",CellIsRule(operator="equal",formula=['"Win"'],fill=PatternFill("solid",bgColor="C6EFCE")))
j.conditional_formatting.add(f"N2:N{N+1}",CellIsRule(operator="equal",formula=['"Loss"'],fill=PatternFill("solid",bgColor="FFC7CE")))

# ---------- Stats
st=wb.create_sheet("Stats")
st.column_dimensions["A"].width=34; st.column_dimensions["B"].width=16; st.column_dimensions["C"].width=60
st["A1"]="Stats (all trades in the Journal)"; st["A1"].font=Font(size=15,bold=True,color=NAVY)
L=f"Journal!$L$2:$L${N+1}"; M=f"Journal!$M$2:$M${N+1}"; Nn=f"Journal!$N$2:$N${N+1}"; Q=f"Journal!$Q$2:$Q${N+1}"
rows=[
("Trades",f'=COUNT({L})','0',"Rows with an entry and exit price"),
("Wins",f'=COUNTIF({Nn},"Win")','0',""),
("Losses",f'=COUNTIF({Nn},"Loss")','0',""),
("Win rate",'=IF(B3=0,"",B4/B3)','0.0%',"Wins divided by trades"),
("Total P&L ($)",f'=SUM({L})','#,##0.00;[Red]-#,##0.00',"After fees"),
("Average win ($)",f'=IFERROR(AVERAGEIF({L},">0"),"")','#,##0.00',""),
("Average loss ($)",f'=IFERROR(AVERAGEIF({L},"<0"),"")','#,##0.00;[Red]-#,##0.00',""),
("Payoff ratio",'=IFERROR(B8/ABS(B9),"")','0.00',"Average win divided by average loss"),
("Expectancy per trade ($)",'=IF(B3=0,"",B7/B3)','#,##0.00;[Red]-#,##0.00',"Average P&L per trade"),
("Expectancy per trade (R)",f'=IFERROR(AVERAGE({M}),"")','0.00"R";[Red]-0.00"R"',"Average R on trades that have planned risk"),
("Profit factor",f'=IFERROR(SUMIF({L},">0")/ABS(SUMIF({L},"<0")),"")','0.00',"Gross wins divided by gross losses (needs at least one loss)"),
("Best trade ($)",f'=IF(B3=0,"",MAX({L}))','#,##0.00',""),
("Worst trade ($)",f'=IF(B3=0,"",MIN({L}))','#,##0.00;[Red]-#,##0.00',""),
("Max drawdown ($)",f'=IF(B3=0,"",MAX({Q}))','#,##0.00',"Largest fall from a running-P&L peak"),
("Average planned risk ($)",f'=IFERROR(AVERAGEIF(Journal!$J$2:$J${N+1},">0"),"")','#,##0.00',""),
("Total fees paid ($)",f'=SUM(Journal!$I$2:$I${N+1})','#,##0.00',""),
]
for i,(a,f,nf,note) in enumerate(rows,start=3):
    st.cell(row=i,column=1,value=a).font=Font(bold=True)
    c=st.cell(row=i,column=2,value=f); c.number_format=nf; c.border=box; c.fill=gry
    st.cell(row=i,column=3,value=note).font=Font(italic=True,color="667085")
# NB: row index map: B3 trades,B4 wins,B5 losses,B6 winrate,B7 total,B8 avgwin,B9 avgloss
# longest streaks via helper columns on Journal (S,T)
st["A21"]="Longest win streak"; st["A22"]="Longest losing streak"
st["A21"].font=st["A22"].font=Font(bold=True)
for r in range(2,N+2):
    j.cell(row=r,column=19,value=f'=IF(N{r}="Win",N(S{r-1})+1,0)' if r>2 else f'=IF(N2="Win",1,0)')
    j.cell(row=r,column=20,value=f'=IF(N{r}="Loss",N(T{r-1})+1,0)' if r>2 else f'=IF(N2="Loss",1,0)')
j["S1"]="win streak (helper)"; j["T1"]="loss streak (helper)"
j.column_dimensions["S"].hidden=True; j.column_dimensions["T"].hidden=True
st["B21"]=f"=MAX(Journal!$S$2:$S${N+1})"; st["B22"]=f"=MAX(Journal!$T$2:$T${N+1})"
for a in ("B21","B22"): st[a].border=box; st[a].fill=gry
# equity chart
ch=LineChart(); ch.title="Running P&L ($)"; ch.height=8; ch.width=18
ch.add_data(Reference(j,min_col=15,min_row=1,max_row=60),titles_from_data=True); ch.legend=None
st.add_chart(ch,"E3")

# ---------- By Setup
b=wb.create_sheet("By Setup")
for i,w in enumerate([22,9,9,12,14,14],start=1): b.column_dimensions[chr(64+i)].width=w
for i,h in enumerate(["Setup tag","Trades","Win rate","Total P&L ($)","Expectancy ($)","Expectancy (R)"],start=1):
    c=b.cell(row=1,column=i,value=h); c.font=hdr; c.fill=hfill; c.alignment=Alignment(horizontal="center",wrap_text=True)
b["H1"]="Type your setup tags in column A (same spelling as the Journal). Sample tags are filled in."; b["H1"].font=Font(italic=True)
C=f"Journal!$C$2:$C${N+1}"
for r,tag in enumerate(["ORB","Reversal","Vertical","","","","","","",""],start=2):
    c=b.cell(row=r,column=1,value=tag or None); c.fill=inp; c.border=box
    b.cell(row=r,column=2,value=f'=IF(A{r}="","",COUNTIFS({C},A{r},{L},"<>"))')
    b.cell(row=r,column=3,value=f'=IF(OR(A{r}="",N(B{r})=0),"",COUNTIFS({C},A{r},{Nn},"Win")/B{r})').number_format="0.0%"
    b.cell(row=r,column=4,value=f'=IF(A{r}="","",SUMIFS({L},{C},A{r}))').number_format='#,##0.00;[Red]-#,##0.00'
    b.cell(row=r,column=5,value=f'=IF(OR(A{r}="",N(B{r})=0),"",D{r}/B{r})').number_format='#,##0.00;[Red]-#,##0.00'
    b.cell(row=r,column=6,value=f'=IF(OR(A{r}="",N(B{r})=0),"",IFERROR(AVERAGEIFS({M},{C},A{r}),""))').number_format='0.00"R"'
    for c in range(2,7): b.cell(row=r,column=c).fill=gry; b.cell(row=r,column=c).border=box

# ---------- Calendar
k=wb.create_sheet("Calendar")
k["A1"]="Month to show (type the first day of the month):"; k["A1"].font=Font(bold=True)
k["E1"]=dt.date(2026,10,1); k["E1"].number_format="mmmm yyyy"; k["E1"].fill=inp; k["E1"].border=box
for i,d in enumerate(["Mon","Tue","Wed","Thu","Fri"]):
    c=k.cell(row=3,column=2+i,value=d); c.font=hdr; c.fill=hfill; c.alignment=Alignment(horizontal="center")
k.cell(row=3,column=1,value="Week of").font=hdr; k.cell(row=3,column=1).fill=hfill
k.cell(row=3,column=7,value="Week total").font=hdr; k.cell(row=3,column=7).fill=hfill
k.column_dimensions["A"].width=14
for col in "BCDEFG": k.column_dimensions[col].width=14
D=f"Journal!$A$2:$A${N+1}"
# Monday of week containing 1st: $E$1 - WEEKDAY($E$1,3)
for w in range(6):
    r=4+w*2
    k.cell(row=r,column=1,value=f"=$E$1-WEEKDAY($E$1,3)+{7*w}").number_format="d mmm"
    k.cell(row=r,column=1).font=Font(bold=True)
    for d in range(5):
        col=2+d
        dayf=f"($A{r}+{d})"
        k.cell(row=r,column=col,value=f'=IF(MONTH({dayf})=MONTH($E$1),DAY({dayf}),"")').font=Font(size=9,color="667085")
        c=k.cell(row=r+1,column=col,value=f'=IF(MONTH({dayf})<>MONTH($E$1),"",IF(COUNTIF({D},{dayf})=0,"",SUMIFS({L},{D},{dayf})))')
        c.number_format='#,##0.00;[Red]-#,##0.00'; c.border=box; c.alignment=Alignment(horizontal="center")
    g=k.cell(row=r+1,column=7,value=f'=IF(COUNT(B{r+1}:F{r+1})=0,"",SUM(B{r+1}:F{r+1}))'); g.number_format='#,##0.00;[Red]-#,##0.00'; g.font=Font(bold=True); g.border=box
k["A17"]="Month total ($)"; k["A17"].font=Font(bold=True)
k["B17"]="=SUM(G5,G7,G9,G11,G13,G15)"; k["B17"].number_format='#,##0.00;[Red]-#,##0.00'; k["B17"].font=Font(bold=True)
k["A18"]="Green days / red days"; k["A18"].font=Font(bold=True)
rng="B5:F5,B7:F7,B9:F9,B11:F11,B13:F13,B15:F15"
k["B18"]='=COUNTIF(B5:F5,">0")+COUNTIF(B7:F7,">0")+COUNTIF(B9:F9,">0")+COUNTIF(B11:F11,">0")+COUNTIF(B13:F13,">0")+COUNTIF(B15:F15,">0")'
k["C18"]='=COUNTIF(B5:F5,"<0")+COUNTIF(B7:F7,"<0")+COUNTIF(B9:F9,"<0")+COUNTIF(B11:F11,"<0")+COUNTIF(B13:F13,"<0")+COUNTIF(B15:F15,"<0")'
k["A19"]="Weekend trades are not shown (US options trade Mon-Fri)."; k["A19"].font=Font(italic=True,color="667085")
for ref in ["B5:F5","B7:F7","B9:F9","B11:F11","B13:F13","B15:F15"]:
    k.conditional_formatting.add(ref,CellIsRule(operator="greaterThan",formula=["0"],fill=PatternFill("solid",bgColor="C6EFCE")))
    k.conditional_formatting.add(ref,CellIsRule(operator="lessThan",formula=["0"],fill=PatternFill("solid",bgColor="FFC7CE")))
wb.save("../Options Day Trader Journal.xlsx")
print("saved")
