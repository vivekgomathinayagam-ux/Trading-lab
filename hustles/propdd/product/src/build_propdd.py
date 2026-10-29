from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
NAVY="1F2A44"; inp=PatternFill("solid",fgColor="FFF8DC"); gry=PatternFill("solid",fgColor="F2F4F7")
hdr=Font(bold=True,color="FFFFFF"); hfill=PatternFill("solid",fgColor=NAVY)
thin=Side(style="thin",color="D0D5DD"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
USD='$#,##0.00'; PCT='0.0%'
wb=Workbook(); s=wb.active; s.title="Start here"
s.column_dimensions["A"].width=3; s.column_dimensions["B"].width=100
lines=[("Funded-Account Drawdown Tracker",18,True),
("Track a trading-challenge or funded account against its own loss rules, one row per trading day.",11,False),("",11,False),
("How to use it",13,True),
("1. 'Rules' tab: type the numbers from YOUR firm's rules (account size, profit target, daily loss limit, max drawdown, trailing or static). Yellow cells are inputs.",11,False),
("2. 'Daily log' tab: add one row per trading day with the day's net profit or loss in dollars (after fees). Everything else calculates.",11,False),
("3. 'Rules' tab also shows your current balance, your drawdown floor, how many dollars of cushion are left, and a status line.",11,False),
("4. Rows 5-12 of the Daily log hold SAMPLE numbers so you can see how it works. Delete the yellow P&L values to start your own.",11,False),("",11,False),
("Read this first",13,True),
("Firms differ: some trail the floor on end-of-day balance, some on intraday equity, some lock the floor at the starting balance, some measure daily loss from the day's starting balance or from equity. This sheet models the common end-of-day versions only. Always check your firm's own rules and dashboard; they are the authority, this sheet is a second set of eyes.",11,False),
("This is arithmetic on numbers you type. It is not advice, not a signal, and not affiliated with any firm. Trading involves risk of loss. Built with AI assistance (Claude) and tested by recalculating every formula.",11,False)]
for i,(t,z,b) in enumerate(lines,start=2):
    c=s.cell(row=i,column=2,value=t); c.font=Font(size=z,bold=b,color=NAVY if b else "000000"); c.alignment=Alignment(wrap_text=True,vertical="top")
r=wb.create_sheet("Rules"); r.column_dimensions["A"].width=40; r.column_dimensions["B"].width=18; r.column_dimensions["C"].width=70
r["A1"]="Account rules and live status"; r["A1"].font=Font(size=15,bold=True,color=NAVY)
def put(row,label,val,nf=None,i=False,note=""):
    r.cell(row=row,column=1,value=label).font=Font(bold=True)
    c=r.cell(row=row,column=2,value=val); c.border=box; c.fill=inp if i else gry
    if nf: c.number_format=nf
    r.cell(row=row,column=3,value=note).font=Font(italic=True,color="667085")
put(3,"Starting balance ($)",50000,USD,True,"Account size at the start of the challenge or funded stage")
put(4,"Profit target ($)",3000,USD,True,"Dollars of profit needed (0 if none, e.g. funded stage)")
put(5,"Max daily loss ($)",1000,USD,True,"Loss in one day that breaches the account (0 if none)")
put(6,"Max drawdown ($)",2000,USD,True,"Distance from the high-water mark (trailing) or from the start (static)")
put(7,"Drawdown type","Trailing (end of day)",None,True,"Trailing: floor rises with your best end-of-day balance. Static: floor never moves")
put(8,"Floor stops rising at starting balance?","Yes",None,True,"Many firms stop trailing once the floor reaches the starting balance. Pick what your firm does")
dv1=DataValidation(type="list",formula1='"Trailing (end of day),Static"',allow_blank=False); dv2=DataValidation(type="list",formula1='"Yes,No"',allow_blank=False)
r.add_data_validation(dv1); r.add_data_validation(dv2); dv1.add("B7"); dv2.add("B8")
r["A10"]="Live status (calculated)"; r["A10"].font=Font(bold=True,size=12,color=NAVY)
N=200; last=4+N  # log rows 5..204
rng=lambda col:f"'Daily log'!{col}5:{col}{last}"
put(11,"Trading days logged",f"=COUNT({rng('B')})","0")
put(12,"Current balance ($)",f"=B3+SUM({rng('B')})",USD)
put(13,"Highest end-of-day balance ($)",f"=MAX(B3,MAX({rng('D')}))",USD)
put(14,"Drawdown floor now ($)",f"=IF(B7=\"Static\",B3-B6,IF(B8=\"Yes\",MIN(B3,B13-B6),B13-B6))",USD,note="Balance at or below this = account breached")
put(15,"Cushion above the floor ($)","=B12-B14",USD,note="How much you can still lose before the drawdown rule is breached")
put(16,"Cushion as % of balance","=IF(B12=0,0,B15/B12)",PCT)
put(17,"Last day's P&L ($)",f"=IF(B11=0,0,INDEX({rng('B')},B11))",USD)
put(18,"Daily loss limit used (last day)","=IF(OR(B5=0,B17>=0),0,-B17/B5)",PCT,note="100% or more = daily limit breached")
put(19,"Profit still needed to hit target ($)","=MAX(0,B4-(B12-B3))",USD)
put(20,"Progress to target","=IF(B4=0,0,MAX(0,B12-B3)/B4)",PCT)
put(21,"Days with a breach",f"=COUNTIF({rng('J')},\"BREACH*\")","0")
put(22,"Status",'=IF(B11=0,"No days logged",IF(B21>0,"BREACHED: check your firm",IF(B12<=B14,"BREACHED: at the floor",IF(AND(B4>0,B12-B3>=B4),"Target reached",IF(B15<B6*0.25,"CAUTION: under 25% of drawdown left","OK")))))',None,note="Re-check against your firm's own dashboard")
r["B22"].font=Font(bold=True)
r.conditional_formatting.add("B22",FormulaRule(formula=['LEFT(B22,8)="BREACHED"'],fill=PatternFill("solid",bgColor="F8B4B4")))
r.conditional_formatting.add("B22",FormulaRule(formula=['LEFT(B22,7)="CAUTION"'],fill=PatternFill("solid",bgColor="FDE68A")))
r.conditional_formatting.add("B22",FormulaRule(formula=['OR(B22="OK",B22="Target reached")'],fill=PatternFill("solid",bgColor="BBF7D0")))
d=wb.create_sheet("Daily log")
heads=["Date","Net P&L ($)","End balance ($)","Highest end balance ($)","Drawdown floor ($)","Cushion ($)","Daily loss used","Profit to date ($)","Note","Check"]
w=[12,14,16,18,16,14,14,16,28,22]
for i,(h,x) in enumerate(zip(heads,w),start=1):
    c=d.cell(row=4,column=i,value=h); c.font=hdr; c.fill=hfill; c.alignment=Alignment(wrap_text=True,horizontal="center"); d.column_dimensions[chr(64+i)].width=x
d["A1"]="Daily log: type a date and the day's net P&L; the grey columns calculate"; d["A1"].font=Font(size=14,bold=True,color=NAVY)
d["A2"]="Rows 5-12 are SAMPLE data. Clear the yellow cells (A and B, and I if used) to start your own."; d["A2"].font=Font(italic=True,color="B42318")
sample=[("2026-09-01",420),("2026-09-02",-310),("2026-09-03",880),("2026-09-04",-850),("2026-09-07",650),("2026-09-08",1200),("2026-09-09",-450),("2026-09-10",300)]
import datetime
for k in range(N):
    row=5+k; p=row-1
    if k<len(sample):
        d.cell(row=row,column=1,value=datetime.date.fromisoformat(sample[k][0])).number_format="yyyy-mm-dd"; d.cell(row=row,column=2,value=sample[k][1])
        d.cell(row=row,column=9,value="SAMPLE")
    else: d.cell(row=row,column=1).number_format="yyyy-mm-dd"
    for col in (1,2,9): d.cell(row=row,column=col).fill=inp; d.cell(row=row,column=col).border=box
    d.cell(row=row,column=2).number_format=USD
    d.cell(row=row,column=3,value=f'=IF(B{row}="","",Rules!$B$3+SUM($B$5:B{row}))')
    d.cell(row=row,column=4,value=f'=IF(B{row}="","",MAX(Rules!$B$3,MAX($C$5:C{row})))')
    d.cell(row=row,column=5,value=f'=IF(B{row}="","",IF(Rules!$B$7="Static",Rules!$B$3-Rules!$B$6,IF(Rules!$B$8="Yes",MIN(Rules!$B$3,D{row}-Rules!$B$6),D{row}-Rules!$B$6)))')
    d.cell(row=row,column=6,value=f'=IF(B{row}="","",C{row}-E{row})')
    d.cell(row=row,column=7,value=f'=IF(B{row}="","",IF(OR(Rules!$B$5=0,B{row}>=0),0,-B{row}/Rules!$B$5))').number_format=PCT
    d.cell(row=row,column=8,value=f'=IF(B{row}="","",C{row}-Rules!$B$3)')
    d.cell(row=row,column=10,value=f'=IF(B{row}="","",IF(C{row}<=E{row},"BREACH: floor",IF(G{row}>=1,"BREACH: daily limit",IF(F{row}<Rules!$B$6*0.25,"Caution","OK"))))')
    for col in (3,4,5,6,8): d.cell(row=row,column=col).number_format=USD; d.cell(row=row,column=col).fill=gry
    d.cell(row=row,column=7).fill=gry; d.cell(row=row,column=10).fill=gry
d.freeze_panes="A5"
d.conditional_formatting.add(f"J5:J{last}",FormulaRule(formula=['LEFT(J5,6)="BREACH"'],fill=PatternFill("solid",bgColor="F8B4B4")))
wb.move_sheet("Rules",offset=0)
wb.save("Funded Account Drawdown Tracker.xlsx")
